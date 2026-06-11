# -*- coding: utf-8 -*-
import zipfile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from .config import input_dir, output_dir, template_dir
from .parser import load_kegiatan_tambahan, parse_detail_github, parse_weekly_report

def clear_body(doc):
    body = doc.element.body
    for child in list(body):
        if child.tag.split("}")[-1] != "sectPr":
            body.remove(child)

def enable_update_fields(doc):
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    doc.settings.element.append(update)

def add_text(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Calibri"
    r.font.size = Pt(11)

def add_center(doc, text, bold=False, size=11):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = "Calibri"

def add_heading(doc, text, level=1):
    doc.add_heading(text, level=level)

def add_toc(doc):
    add_heading(doc, "DAFTAR ISI", 1)
    p = doc.add_paragraph()
    r = p.add_run()
    for el_type, text in [("begin", None), ("instr", ' TOC \\o "1-3" \\h \\z \\u '), ("separate", None), ("end", None)]:
        if el_type == "instr":
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = text
            r._r.append(instr)
        else:
            fld = OxmlElement("w:fldChar")
            fld.set(qn("w:fldCharType"), el_type)
            r._r.append(fld)
    add_text(doc, "(Klik kanan Daftar Isi -> Update Field jika nomor halaman belum muncul)")

def add_activity(doc, activity):
    add_text(doc, activity["subtitle"], bold=True)
    if activity.get("prolog"):
        add_text(doc, "Prolog")
        add_text(doc, activity["prolog"])
    if activity.get("deskripsi"):
        add_text(doc, "Deskripsi Pekerjaan")
        add_text(doc, activity["deskripsi"])
    add_text(doc, "Detail Perubahan")
    if activity.get("files"):
        add_text(doc, "File yang Diubah:")
        for f in activity["files"]:
            add_text(doc, f)
    if activity.get("utama"):
        add_text(doc, "Perubahan Utama:")
        for u in activity["utama"]:
            add_text(doc, u)
    if activity.get("manfaat"):
        add_text(doc, "Manfaat:")
        for m in activity["manfaat"]:
            add_text(doc, m)
    add_text(doc, "Dokumentasi")
    for d in activity.get("docs") or []:
        add_text(doc, d)

def ensure_media(media_dir, template_path):
    if media_dir.exists() and any(media_dir.iterdir()):
        return
    media_dir.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(template_path) as z:
        for name in z.namelist():
            if name.startswith("word/media/"):
                (media_dir / name.split("/")[-1]).write_bytes(z.read(name))

def add_image_if_exists(doc, media_dir, name, width=Inches(6.0)):
    path = media_dir / name
    if path.exists():
        doc.add_picture(str(path), width=width)
        return True
    return False

def add_pr_table(doc, prs):
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for i, h in enumerate(["No", "PR #", "Judul", "Status"]):
        table.rows[0].cells[i].text = h
    for idx, pr in enumerate(prs, 1):
        row = table.add_row().cells
        row[0].text = str(idx)
        row[1].text = str(pr["number"])
        row[2].text = pr["title"]
        row[3].text = pr["state"].upper()

def build_bab1(doc, config):
    bab1 = config.get("bab1", {})
    bulan, tahun = config["bulan"], config["tahun"]
    add_heading(doc, "BAB I", 1)
    add_heading(doc, "PENDAHULUAN", 1)
    add_text(doc, "Latar Belakang", bold=True)
    add_text(doc, bab1.get("latar_belakang_pembuka", f"Laporan kemajuan pekerjaan bulan {bulan} {tahun}."))
    add_text(doc, bab1.get("latar_belakang_sistem", "Latar belakang implementasi logika sistem pengawasan SDKP."))
    if bab1.get("pencapaian"):
        add_text(doc, bab1["pencapaian"])
    if bab1.get("tantangan"):
        add_text(doc, bab1["tantangan"])
    add_text(doc, "Maksud dan Tujuan", bold=True)
    add_text(doc, bab1.get("maksud_tujuan", "Memberikan pemahaman transparan mengenai kemajuan pekerjaan."))
    add_text(doc, "Lingkup Pekerjaan", bold=True)
    for item in bab1.get("lingkup", ["Merancang Arsitektur Backend", "Pemrograman dan Pengembangan", "Manajemen Basis Data", "Manajemen GitHub", "Pembaruan API Server Development", "Koordinasi Tim Frontend", "Manajemen Error dan Logging"]):
        add_text(doc, item)

def build_bab2(doc, modules):
    add_heading(doc, "BAB II", 1)
    add_heading(doc, "HASIL KEGIATAN", 1)
    for mod in modules:
        if mod["title"].upper().startswith("RINGKASAN"):
            continue
        add_heading(doc, mod["title"], 2)
        if mod.get("intro"):
            add_text(doc, mod["intro"])
        for act in mod["activities"]:
            add_activity(doc, act)

def build_kegiatan_tambahan(doc, kegiatan):
    if not kegiatan:
        return
    add_heading(doc, "Kegiatan Tambahan", 2)
    for item in kegiatan:
        add_heading(doc, item["judul"], 3)
        add_text(doc, item.get("deskripsi", ""))
        if item.get("link"):
            add_text(doc, item["link"])

def build_manajemen_kode(doc, config, prs, media_dir):
    bulan, tahun = config["bulan"], config["tahun"]
    gh, media = config["github"], config.get("media", {})
    add_heading(doc, "Manajemen Kode Sumber", 2)
    add_text(doc, f"Singkronisasi Kode Sumber pada Repositori Github {bulan} {tahun}")
    add_text(doc, f"Pada bulan {bulan} {tahun} dilakukan {len(prs)} pull request (author: {gh['author']}) di https://github.com/{gh['repo']}")
    add_text(doc, "Gambar Tangkapan Layar Terkait Proses Sinkronisasi Kode Sumber di Github")
    for img in media.get("github_screenshots", [f"image{i}.png" for i in range(1, 11)]):
        add_image_if_exists(doc, media_dir, img)
    add_text(doc, f"Singkronisasi Kode Sumber pada Server Development {bulan} {tahun}")
    add_text(doc, "Sinkronisasi kode sumber API ke server development setelah perubahan di-merge.")
    add_text(doc, "Gambar Tangkapan Layar Terkait Proses Integrasi API di Server PSDKP")
    for img in media.get("server_screenshots", ["image11.jpeg", "image12.png"]):
        add_image_if_exists(doc, media_dir, img)

def build_bab3(doc, config, prs):
    bab3, bulan = config.get("bab3", {}), config["bulan"]
    add_heading(doc, "BAB III", 1)
    add_heading(doc, "PENUTUP", 1)
    add_text(doc, "Kesimpulan", bold=True)
    add_text(doc, bab3.get("kesimpulan_pembuka", f"Kemajuan implementasi logika sistem pada bulan {bulan} {config['tahun']}."))
    for item in bab3.get("kesimpulan_items", [f"Manajemen GitHub   {len(prs)} pull request pada bulan {bulan}"]):
        add_text(doc, item)
    add_text(doc, "Saran", bold=True)
    for s in bab3.get("saran_items", ["Peningkatan Keamanan Sistem", "Pengalaman Pengguna", "Sistem Pemantauan via Sentry", "Evaluasi Feedback Pengguna"]):
        add_text(doc, s)

def build_lampiran(doc, config, prs, weekly_rows, media_dir):
    bulan, tech = config["bulan"], config.get("tech", {})
    snippet, media = config.get("code_snippet", {}), config.get("media", {})
    doc.add_page_break()
    add_heading(doc, "LAMPIRAN", 1)
    add_text(doc, f"Lampiran 1 Weekly Report {bulan} {config['tahun']}", bold=True)
    if weekly_rows:
        tbl = doc.add_table(rows=1, cols=5)
        tbl.style = "Table Grid"
        for i, h in enumerate(["Minggu", "Tanggal", "Hari", "Aktivitas", "Dokumentasi"]):
            tbl.rows[0].cells[i].text = h
        for row in weekly_rows:
            cells = tbl.add_row().cells
            cells[0].text = row.get("minggu", "")
            cells[1].text = row.get("tanggal", "")
            cells[2].text = row.get("hari", "")
            cells[3].text = row.get("aktivitas", "")
            cells[4].text = row.get("dokumentasi", "")
    doc.add_paragraph()
    add_text(doc, "Lampiran 2. Kode Sumber", bold=True)
    add_text(doc, f"Bahasa Pemrograman : {tech.get('language', 'PHP v8.2')}")
    add_text(doc, f"Framework : {tech.get('framework', 'Laravel 11')}")
    add_text(doc, f"Repositori : {tech.get('repo_url', 'https://github.com/setditjen-psdkp/api-sip')}")
    add_text(doc, f"Daftar Pull Request Bulan {bulan} {config['tahun']} (Author: {config['github']['author']}) - Total: {len(prs)} PR")
    add_pr_table(doc, prs)
    doc.add_paragraph()
    add_text(doc, "Gambar Tangkapan Layar Repositori Github dan Pull Request")
    for img in media.get("lampiran_screenshots", ["image13.jpeg", "image14.jpeg", "image15.jpeg", "image16.jpeg"]):
        add_image_if_exists(doc, media_dir, img)
    if snippet.get("content"):
        add_text(doc, snippet.get("title", "Contoh Cuplikan Kode Sumber"))
        p = doc.add_paragraph()
        r = p.add_run(snippet["content"])
        r.font.name = "Courier New"
        r.font.size = Pt(9)
    add_text(doc, f"Lampiran 3 Detail Laporan GitHub {bulan} {config['tahun']}", bold=True)
    add_text(doc, f"File DOCX: {config.get('detail_docx', '')}")
    if config.get("detail_txt"):
        add_text(doc, f"File TXT: {config.get('detail_txt', '')}")

def build_report(config, prs):
    period = config["period"]
    inp, tpl = input_dir(period), template_dir()
    media_dir, template_path = tpl / "media", tpl / "laporan_template.docx"
    output_path = output_dir(period) / config["output_filename"]
    ensure_media(media_dir, template_path)
    modules = parse_detail_github(inp / "detail_github.md") if (inp / "detail_github.md").exists() else []
    weekly_rows = parse_weekly_report(inp / "weekly_report.md", config.get("weekly_user_filter", "Lutfi"))
    kegiatan = load_kegiatan_tambahan(inp / "kegiatan_tambahan.json")
    doc = Document(str(template_path))
    clear_body(doc)
    enable_update_fields(doc)
    bulan_up = config.get("bulan_up", config["bulan"].upper())
    nama = config.get("nama", "Lutfi Ihsan")
    for _ in range(3):
        doc.add_paragraph()
    add_center(doc, "LAPORAN KEMAJUAN", True, 14)
    add_center(doc, "TENAGA TEKNIS IMPLEMENTASI LOGIKA SISTEM", True, 14)
    add_center(doc, "DALAM RANGKA PENGELOLAAN DATA PENGAWASAN SDKP", True, 12)
    add_center(doc, f"BULAN {bulan_up} {config['tahun']}", True, 12)
    doc.add_paragraph()
    add_center(doc, nama)
    add_center(doc, config.get("jabatan", "Tenaga Teknis Implementasi Logika Sistem"))
    doc.add_paragraph()
    add_center(doc, "DIREKTORAT JENDERAL PENGAWASAN SUMBER DAYA")
    add_center(doc, "KELAUTAN DAN PERIKANAN")
    add_center(doc, "KEMENTERIAN KELAUTAN DAN PERIKANAN")
    add_center(doc, str(config["tahun"]))
    doc.add_page_break()
    add_toc(doc)
    doc.add_page_break()
    add_heading(doc, "Kata Pengantar", 1)
    add_text(doc, config.get("kata_pengantar_1", f"Laporan kemajuan bulan {config['bulan']} {config['tahun']}."))
    add_text(doc, config.get("kata_pengantar_2", "Memantau, menyusun, mengimplementasikan, dan memelihara sistem pengawasan SDKP."))
    add_text(doc, f"Jakarta,    {config['bulan']} {config['tahun']}")
    add_text(doc, "Hormat Kami,")
    doc.add_paragraph()
    add_text(doc, nama)
    doc.add_page_break()
    build_bab1(doc, config)
    doc.add_page_break()
    build_bab2(doc, modules)
    doc.add_page_break()
    build_kegiatan_tambahan(doc, kegiatan)
    build_manajemen_kode(doc, config, prs, media_dir)
    doc.add_page_break()
    build_bab3(doc, config, prs)
    build_lampiran(doc, config, prs, weekly_rows, media_dir)
    doc.save(str(output_path))
    return output_path
